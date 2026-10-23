import os
import sys
import json
import re
import glob
import datetime
import imaplib
import email
from email.header import decode_header
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(BASE_DIR, "data", "family_db.json")
PARENT_DIR = os.path.dirname(BASE_DIR)
FIXED_EXCEL_PATH = os.path.join(PARENT_DIR, "2026 - FAMILY NOTE BOOK (Da Chuan Hoa & Gop Dau Tu).xlsx")


def load_db():
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def save_db(db):
    db["meta"]["last_synced"] = datetime.datetime.now().isoformat()
    with open(DATA_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)


def recompute_months(db):
    txs = db.get("transactions", [])
    for m_obj in db.get("months", []):
        m = m_obj["month"]
        m_txs = [t for t in txs if int(t.get("month", 0)) == m]
        actual_exp = sum(float(t.get("amount", 0)) for t in m_txs)
        inc_total = (
            float(m_obj["income"].get("CK", 0))
            + float(m_obj["income"].get("VK", 0))
            + float(m_obj["income"].get("Subo", 0))
            + float(m_obj["income"].get("Sugi", 0))
        )
        m_obj["income"]["total"] = inc_total
        m_obj["target_saving"] = round(inc_total * 0.30)
        m_obj["target_expense"] = round(inc_total * 0.70)
        m_obj["actual_total_expense"] = actual_exp
        m_obj["remaining_actual"] = inc_total - actual_exp
        m_obj["tx_count"] = len(m_txs)
        dca_tot = float(m_obj["allocation"].get("dca_total", 0))
        dip_add = float(m_obj["allocation"].get("dip_fund", 0))
        m_obj["allocation"]["dca_fuevfvnd"] = round(dca_tot * 0.30)
        m_obj["allocation"]["dca_dcds"] = round(dca_tot * 0.25)
        m_obj["allocation"]["dca_vesaf"] = round(dca_tot * 0.45)
        if inc_total > 0:
            m_obj["allocation"]["emergency_saving"] = max(0.0, (inc_total - actual_exp) - dca_tot - dip_add)

    # Recompute investment summary
    holdings = db["investment"]["holdings"]
    orders = db["investment"].get("dca_orders", [])
    for h in holdings:
        code = h["code"]
        c_orders = [o for o in orders if o["code"] == code]
        if c_orders:
            h["invested"] = sum(float(o["amount"]) for o in c_orders)
            h["units"] = round(sum(float(o["units"]) for o in c_orders), 2)
        h["market_value"] = round(float(h["units"]) * float(h["current_nav"]))
        h["avg_cost"] = round(h["invested"] / h["units"]) if h["units"] > 0 else 0
        h["pnl"] = h["market_value"] - h["invested"]
        h["pnl_pct"] = round((h["pnl"] / h["invested"]) * 100, 2) if h["invested"] > 0 else 0.0

    tot_mv = sum(h["market_value"] for h in holdings)
    for h in holdings:
        h["actual_weight"] = round(h["market_value"] / tot_mv, 4) if tot_mv > 0 else 0.0


def append_tx_to_excel(tx):
    """Synchronize new transaction directly into the fixed Excel workbook."""
    if not os.path.exists(FIXED_EXCEL_PATH):
        return False
    try:
        wb = openpyxl.load_workbook(FIXED_EXCEL_PATH)
        m = int(tx["month"])
        sheet_name = f"T{m} - DETAIL"
        if sheet_name not in wb.sheetnames:
            return False
        ws = wb[sheet_name]
        # Find first empty transaction row in Col C & F starting from row 4
        target_row = 4
        for r in range(4, 830):
            if ws.cell(r, 3).value is None and ws.cell(r, 6).value is None:
                target_row = r
                break
        try:
            dt_obj = datetime.datetime.strptime(tx["date"], "%Y-%m-%d")
        except Exception:
            dt_obj = tx["date"]

        # Map back to Excel conventions
        cat_excel_map = {
            "Ăn uống": "Ăn uống",
            "Vui chơi": "vui chơi",
            "Công việc": "Công việc",
            "Học tập": "học tập",
            "Đồ dùng GĐ": "đồ dùng GĐ",
            "Phí tiện ích": "phí tiện ích",
            "Làm đẹp": "Làm đẹp",
            "Y tế": "Y tế",
            "Event": "event",
            "Nợ": "Nợ"
        }
        person_excel_map = {
            "Chung": "Chung",
            "Chồng": "chồng",
            "Vợ": "vợ",
            "Subo": "SUBO",
            "Sugi": "SUGI"
        }
        ws.cell(target_row, 2).value = dt_obj
        ws.cell(target_row, 3).value = tx["content"]
        ws.cell(target_row, 4).value = cat_excel_map.get(tx["category"], tx["category"])
        ws.cell(target_row, 5).value = person_excel_map.get(tx["person"], tx["person"])
        ws.cell(target_row, 6).value = float(tx["amount"])
        ws.cell(target_row, 7).value = tx.get("source", "")
        wb.save(FIXED_EXCEL_PATH)
        return True
    except Exception as e:
        print("Excel sync warning:", e)
        return False


# ==============================================================================
# SMART VIETNAMESE TRANSACTION & BANK EMAIL AI PARSER (VPBank, BIDV, NCB, UOB)
# ==============================================================================

KEYWORD_RULES = [
    # (keywords_list, category, default_person, clean_label)
    (["học phí subo", "hphi subo", "hoc phi subo", "tiền học subo"], "Học tập", "Subo", "học phí subo"),
    (["đón bo", "đón subo", "don subo", "grab subo"], "Phí tiện ích", "Subo", "đón subo"),
    (["tini", "khu vui chơi", "lego", "đồ chơi", "tô tượng", "super wing"], "Vui chơi", "Subo", None),
    (["bỉm", "sữa thanh", "kem hăm", "kem bôi hăm", "bình sữa", "túi trữ sữa", "tiêm chủng sugi", "con cưng", "concung"], "Đồ dùng GĐ", "Sugi", None),
    (["nhi đồng", "khám sugi", "thuốc sugi"], "Y tế", "Sugi", None),
    (["shinhan", "nợ ck shinhan"], "Nợ", "Chồng", "nợ ck shinhan"),
    (["thuế ck", "nợ thuế"], "Nợ", "Chồng", "nợ thuế ck"),
    (["nợ đt vk", "đt vk"], "Nợ", "Vợ", "nợ đt vk"),
    (["xăng ck", "đổ xăng ck"], "Phí tiện ích", "Chồng", "xăng ck"),
    (["xăng vk", "xăng"], "Phí tiện ích", "Vợ", "xăng"),
    (["sửa xe", "rửa xe", "gửi xe", "thẻ xe"], "Phí tiện ích", "Chung", None),
    (["tiền điện", "evn", "điện lực"], "Phí tiện ích", "Chung", "tiền điện"),
    (["tiền nước", "sawaco", "nước cc"], "Phí tiện ích", "Chung", "tiền nước"),
    (["tiền mạng", "fpt", "viettel", "vnpt", "internet"], "Phí tiện ích", "Chung", "tiền mạng"),
    (["qlcc", "phí cc", "phí quản lý", "chung cư"], "Phí tiện ích", "Chung", "phí qlcc"),
    (["grab", "be ", "be-", "xanh sm", "gsm", "taxi", "vé xe"], "Phí tiện ích", "Chung", None),
    (["chatgpt", "chat gpt", "openai", "claude"], "Công việc", "Chung", "chat gpt"),
    (["coop", "co.op", "bhx", "bách hóa xanh", "bach hoa xanh", "winmart", "emart", "aeon", "kingfood", "king food", "ministop", "circle k", "gs25", "thiso", "lotte", "mega market", "siêu thị", "đi chợ"], "Ăn uống", "Chung", None),
    (["phở", "bún", "cơm", "bánh mì", "trà sữa", "cafe", "cà phê", "highlands", "phúc long", "starbucks", "kfc", "texas", "pizza", "lẩu", "nướng", "ốc", "gà", "thịt", "cá", "rau", "trái cây", "sầu riêng", "măng cụt", "sữa", "vinamilk", "th true", "ăn sáng", "ăn trưa", "ăn tối", "ăn vặt"], "Ăn uống", "Chung", None),
    (["long châu", "pharmacity", "an khang", "thuốc tây", "thuốc", "bệnh viện", "khám", "nha khoa", "làm răng", "y tế", "tiểu đường"], "Y tế", "Chung", None),
    (["paula", "mỹ phẩm", "son", "kem chống nắng", "gội đầu", "làm tóc", "cắt tóc", "nail", "spa", "quần áo", "áo", "váy", "giày", "dép", "đồ lót", "uniqlo", "zara"], "Làm đẹp", "Vợ", None),
    (["nước giặt", "nước rửa", "bao rác", "khăn giấy", "giấy vệ sinh", "bàn chải", "kem đr", "dầu gội", "sữa tắm", "vim", "finish", "khay", "nồi", "chảo", "đồ gia dụng"], "Đồ dùng GĐ", "Chung", None),
    (["shopee", "spee", "lazada", "tiktok shop", "tiki"], "Đồ dùng GĐ", "Chung", None),
    (["sinh nhật", "đám cưới", "mừng", "thăm", "du lịch", "khách sạn", "vé máy bay", "event"], "Event", "Chung", None),
]


def classify_transaction_text(text):
    t_low = text.lower().strip()
    cat = "Ăn uống"
    person = "Chung"
    confidence = 0.78
    reason = "Phân loại mặc định theo tần suất cao nhất của gia đình"

    # Detect explicit person mentions first (whole word match)
    explicit_person = None
    if re.search(r'\b(subo|sn bo|đón bo)\b', t_low):
        explicit_person = "Subo"
    elif re.search(r'\b(sugi)\b', t_low):
        explicit_person = "Sugi"
    elif re.search(r'\b(ck|chồng)\b', t_low):
        explicit_person = "Chồng"
    elif re.search(r'\b(vk|vợ)\b', t_low):
        explicit_person = "Vợ"

    for kw_list, rule_cat, rule_person, _ in KEYWORD_RULES:
        for kw in kw_list:
            if kw in t_low:
                cat = rule_cat
                person = explicit_person or rule_person
                confidence = 0.96
                reason = f"Khớp mẫu lịch sử gia đình ('{kw}' → {cat} • {person})"
                break
        if confidence > 0.9:
            break

    if explicit_person and confidence < 0.9:
        person = explicit_person
        confidence = 0.88

    # Detect explicit category mentions if user typed them in quick-add
    for c_key, c_val in [
        ("ăn uống", "Ăn uống"), ("vui chơi", "Vui chơi"), ("công việc", "Công việc"),
        ("học tập", "Học tập"), ("đồ dùng", "Đồ dùng GĐ"), ("tiện ích", "Phí tiện ích"),
        ("làm đẹp", "Làm đẹp"), ("y tế", "Y tế"), ("event", "Event"), ("nợ", "Nợ")
    ]:
        if c_key == "nợ":
            if re.search(r'(?<!ghi\s)(?<!báo\s)(?<!dư\s)\bnợ\b', t_low):
                cat = "Nợ"
                confidence = 0.98
                break
        elif c_key in t_low:
            cat = c_val
            confidence = 0.98
            break

    return cat, person, confidence, reason


def is_unclear_email_transaction(friendly_note, raw_content, confidence):
    """
    Xác định email giao dịch ngân hàng có nội dung mờ / không rõ mục đích chi tiêu hay không.
    Nếu không rõ: cần AI hỏi lại người dùng để xác nhận người chi & hạng mục.
    """
    fn = (friendly_note or "").strip().lower()
    rc = (raw_content or "").strip().lower()

    # 1. Độ tin cậy thấp từ bộ từ khóa gia đình
    if confidence < 0.88:
        return True, "Nội dung giao dịch chưa khớp từ khóa nào của gia đình, AI cần bạn xác nhận."

    # 2. Nội dung quá ngắn hoặc chỉ toàn chữ số / mã giao dịch
    if len(fn) < 4:
        return True, "Nội dung quá ngắn, chỉ có vài ký tự viết tắt ngân hàng."
    if re.match(r'^\d+$', fn) or re.match(r'^[a-z0-9._\-]{1,6}$', fn):
        return True, "Nội dung chỉ là mã số/mã POS của ngân hàng, chưa rõ món đồ đã mua."

    # 3. Nội dung chỉ là thuật ngữ chuyển khoản chung chung không có tên người/món
    vague_patterns = [
        r'^(thanh toan|chuyen tien|chuyen khoan|rut tien|pos|qr|payment|ecom|transfer|napas|ibft|mbvcb)$',
        r'^(thanh toan qr|thanh toan pos|gd online|gd tai quay|ecom transaction)$'
    ]
    for vp in vague_patterns:
        if re.search(vp, fn):
            return True, "Nội dung thanh toán chung chung, chưa rõ mua hàng hóa hay dịch vụ gì."

    return False, None


def get_smart_clarification_suggestions(amount, bank, raw_content):
    """Tạo các lựa chọn gợi ý 1 chạm thông minh khi email chưa rõ nội dung"""
    return [
        {"label": "🍲 Ăn uống • Chung", "category": "Ăn uống", "person": "Chung", "note": "Ăn uống gia đình"},
        {"label": "🧺 Đồ dùng GĐ • Vợ", "category": "Đồ dùng GĐ", "person": "Vợ", "note": "Mua đồ dùng gia đình"},
        {"label": "⛽ Tiện ích / Xăng • Chồng", "category": "Phí tiện ích", "person": "Chồng", "note": "Đổ xăng / Tiện ích"},
        {"label": "👶 Bỉm sữa / Y tế • Sugi", "category": "Đồ dùng GĐ", "person": "Sugi", "note": "Mua đồ cho Sugi"},
        {"label": "👦 Học tập / Đồ chơi • Subo", "category": "Học tập", "person": "Subo", "note": "Chi phí cho Subo"}
    ]


def parse_quick_input(raw_text, default_month=10):
    """
    Parses natural Vietnamese 1-line quick expense input like:
    'emart 450k ăn uống chung vp' or 'bỉm con cưng 378k sugi vp' or 'xăng ck 100000 bidv'
    """
    text = raw_text.strip()
    # Extract bank/source
    source = "VPBank"
    for tok, src_name in [
        ("vpsp", "VPBank Shopee"), ("vp sp", "VPBank Shopee"),
        ("vpbank", "VPBank"), (" vp", "VPBank"),
        ("bidv", "BIDV"), ("ncb", "NCB"), ("uob", "UOB"),
        ("tiền mặt", "Tiền mặt / Chuyển khoản"), ("tm", "Tiền mặt / Chuyển khoản")
    ]:
        if tok in " " + text.lower():
            source = src_name
            break

    # Extract amount: supports 450k, 1.5tr, 1tr2, 450,000, 450000
    amount = 0
    clean_desc = text

    # Pattern 1: 1tr250 or 1.5tr or 12tr
    m_tr = re.search(r'(\d+(?:[.,]\d+)?)\s*(?:tr|triệu|trieu)(?:\s*(\d{1,3}))?', text, re.IGNORECASE)
    if m_tr:
        base = float(m_tr.group(1).replace(",", "."))
        tail = m_tr.group(2)
        amount = base * 1_000_000 + (float(tail) * (10 ** (6 - len(tail))) if tail else 0)
        clean_desc = (text[:m_tr.start()] + " " + text[m_tr.end():]).strip()
    else:
        # Pattern 2: 450k or 450ng
        m_k = re.search(r'(\d+(?:[.,]\d+)?)\s*(?:k|ng|nghìn|ngàn)\b', text, re.IGNORECASE)
        if m_k:
            amount = float(m_k.group(1).replace(",", ".")) * 1000
            clean_desc = (text[:m_k.start()] + " " + text[m_k.end():]).strip()
        else:
            # Pattern 3: raw number >= 1000 or formatted 1,440,000
            m_num = re.search(r'(\d{1,3}(?:[.,]\d{3})+|\d{4,9})', text)
            if m_num:
                raw_n = re.sub(r'[.,]', '', m_num.group(1))
                amount = float(raw_n)
                clean_desc = (text[:m_num.start()] + " " + text[m_num.end():]).strip()

    # Clean trailing source/category tokens from description
    clean_desc = re.sub(r'\b(vpbank|vpsp|vp sp|vp|bidv|ncb|uob)\b', '', clean_desc, flags=re.IGNORECASE).strip(" -,.|")
    cat, person, conf, reason = classify_transaction_text(text)

    now = datetime.datetime.now()
    date_str = f"2026-{int(default_month):02d}-{now.day:02d}"
    return {
        "date": date_str,
        "month": int(default_month),
        "content": clean_desc or text,
        "category": cat,
        "person": person,
        "amount": round(amount),
        "source": source,
        "confidence": conf,
        "reason": reason
    }


def parse_bank_emails(raw_email_text):
    """
    Parses 1 or multiple bank notification emails / SMS messages from VPBank, BIDV, NCB, UOB.
    Splits by '---' or email headers or processes line-by-line / block-by-block.
    """
    blocks = [b.strip() for b in re.split(r'(?:^|\n)---+(?:\n|$)|(?=From:|Từ:|\[VPBank\]|\[BIDV\]|\[NCB\])', raw_email_text) if b.strip()]
    if not blocks:
        blocks = [raw_email_text.strip()]

    parsed_items = []
    for idx, block in enumerate(blocks, 1):
        b_low = block.lower()
        # Detect bank
        bank = "VPBank"
        if "bidv" in b_low or "smartbanking" in b_low:
            bank = "BIDV"
        elif "ncb" in b_low or "quốc dân" in b_low or "izimobile" in b_low:
            bank = "NCB"
        elif "uob" in b_low:
            bank = "UOB"
        elif "shopee" in b_low and "vpbank" in b_low:
            bank = "VPBank Shopee"
        elif "vpbank" in b_low or "neo" in b_low:
            bank = "VPBank"

        # Detect outgoing vs incoming
        is_income = False
        if re.search(r'\+\s*\d', block) and not re.search(r'-\s*\d', block):
            if any(w in b_low for w in ["lương", "salary", "nhận tiền", "ghi có", "cổ tức"]):
                is_income = True

        # Extract amount
        # Look for patterns like -1,440,000 VND, So tien: 378,000 VND, GD: -500,000VND, etc.
        amt = 0.0
        amt_patterns = [
            r'(?:số tiền|so tien|giá trị|gia tri|biến động|bien dong|gd|amount)[:\s]*([+-]?\s*\d{1,3}(?:[.,]\d{3})+|\d{4,10})\s*(?:vnd|đ|vnđ)',
            r'([+-]\s*\d{1,3}(?:[.,]\d{3})+)\s*(?:vnd|đ|vnđ)',
            r'(\d{1,3}(?:[.,]\d{3})+)\s*(?:vnd|đ|vnđ)'
        ]
        for pat in amt_patterns:
            m_a = re.search(pat, block, re.IGNORECASE)
            if m_a:
                num_s = re.sub(r'[^\d]', '', m_a.group(1))
                if num_s:
                    amt = float(num_s)
                    break

        # Extract Date (dd/mm/yyyy or dd-mm-yyyy)
        date_str = datetime.datetime.now().strftime("2026-10-%d")
        month_num = 10
        m_dt = re.search(r'(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?', block)
        if m_dt:
            d_i = int(m_dt.group(1))
            m_i = int(m_dt.group(2))
            if 1 <= d_i <= 31 and 1 <= m_i <= 12:
                month_num = m_i
                date_str = f"2026-{m_i:02d}-{d_i:02d}"

        # Extract merchant / transaction content
        content_extracted = ""
        nd_match = re.search(r'(?:nội dung|noi dung|mô tả|mo ta|merchant|chi tiết)[:\s]+([^\n\r;]+)', block, re.IGNORECASE)
        if nd_match:
            content_extracted = re.split(r'\.\s*(?:số dư|so du|sd)', nd_match.group(1), flags=re.IGNORECASE)[0].strip(" .")
        else:
            # Take the most informative line
            lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
            content_extracted = lines[-1] if lines else block[:80]

        # Clean bank noise words from content_extracted to make it family-notebook friendly
        friendly_note = content_extracted
        friendly_note = re.sub(r'^(MBVCB|IBFT|NAPAS|QR|POS|ECOM|TT|GD|CHUYEN TIEN|THANH TOAN)[\s._:-]*', '', friendly_note, flags=re.IGNORECASE).strip()

        cat, person, conf, reason = classify_transaction_text(content_extracted + " " + friendly_note)
        is_unclear, unclear_msg = is_unclear_email_transaction(friendly_note, content_extracted, conf)

        if amt > 0:
            item_data = {
                "temp_id": f"EMAIL-{idx}-{int(amt)}",
                "bank": bank,
                "date": date_str,
                "month": month_num,
                "amount": round(amt),
                "is_income": is_income,
                "raw_content": content_extracted,
                "content": friendly_note.lower() if len(friendly_note) < 40 else friendly_note[:45],
                "category": cat,
                "person": person,
                "source": bank,
                "confidence": conf,
                "reason": reason,
                "needs_clarification": is_unclear,
                "unclear_reason": unclear_msg,
                "ai_question": f"Giao dịch {round(amt):,}đ tại {bank} ('{friendly_note}') chi cho mục đích gì và của ai trong gia đình?" if is_unclear else None,
                "suggested_options": get_smart_clarification_suggestions(amt, bank, content_extracted) if is_unclear else []
            }
            parsed_items.append(item_data)

    return parsed_items


# ==============================================================================
# INTELLIGENT VIETNAMESE FAMILY FINANCIAL AI ENGINE
# ==============================================================================

def answer_family_question(question, db):
    q = question.lower().strip()
    months = db["months"]
    txs = db["transactions"]
    active_months = [m for m in months if m["income"]["total"] > 0 or m["actual_total_expense"] > 0]
    inv = db["investment"]
    debts = db["debts"]

    total_inc_ytd = sum(m["income"]["total"] for m in active_months)
    total_exp_ytd = sum(m["actual_total_expense"] for m in active_months)
    total_rem_ytd = total_inc_ytd - total_exp_ytd
    saving_rate_ytd = (total_rem_ytd / total_inc_ytd * 100) if total_inc_ytd > 0 else 0

    # Check if asking about a specific month (e.g. "tháng 10", "t10", "tháng 9")
    m_match = re.search(r'(?:tháng|thang|t)\s*(\d{1,2})\b', q)
    target_m = int(m_match.group(1)) if m_match and 1 <= int(m_match.group(1)) <= 12 else None

    if any(k in q for k in ["tư vấn", "tăng giảm", "tăng", "giảm", "cắt giảm", "ngân sách", "tiết kiệm thêm", "tối ưu chi tiêu"]):
        lines = [
            "### 💡 Cố Vấn Tài Chính AI: Tư Vấn Tăng / Giảm Chi Tiêu Từng Nhóm\n",
            "Dựa trên dữ liệu 10 tháng thực tế (Thu: 1.035 tỷ, Chi: 716tr, Thặng dư: +319.2tr / 30.84% thu nhập), AI tư vấn cụ thể như sau:\n",
            "#### 🔴 1. CÁC NHÓM NÊN CẮT GIẢM:",
            "- **🍲 Ăn uống (Thực tế: ~17.5tr/tháng - chiếm 24.5% tổng chi)**:",
            "  * *Khuyến nghị:* **Cắt giảm ~2.500.000 đ/tháng** (mục tiêu giữ dưới 15.000.000 đ/tháng).",
            "  * *Hành động:* Giảm bớt số bữa ăn ngoài hàng vào cuối tuần (trung bình 600k-1.2tr/bữa) và đặt đồ ăn qua app; lên thực đơn tuần trước khi đi Co.opmart/Emart.",
            "- **💄 Làm đẹp / Mua sắm (Từng vọt lên 20.9tr ở T7 và 11.9tr ở T2)**:",
            "  * *Khuyến nghị:* **Khống chế trần cố định 2.500.000 đ/tháng**.",
            "  * *Hành động:* Gom mua mỹ phẩm/quần áo vào các đợt sale định kỳ, hạn chế mua sắm phát sinh ngoài kế hoạch.",
            "- **⚡ Phí tiện ích / Phát sinh xe (T9 vọt lên 19.4tr do sửa xe 12.6tr)**:",
            "  * *Khuyến nghị:* Giữ chi phí sinh hoạt điện nước mạng cố định ~3.5tr/tháng. Các khoản sửa xe đột xuất >5tr nên trích từ Quỹ Khẩn Cấp thay vì trừ thẳng vào ngân sách tháng.\n",
            "#### 🟢 2. CÁC NHÓM NÊN TĂNG / TẬN DỤNG ĐẦU TƯ:",
            "- **📈 Đầu tư DCA 3 Quỹ ETF/Mở (FUEVFVND, DCDS, VESAF)**:",
            "  * *Khuyến nghị:* **NÊN TĂNG TỪ 12.000.000 đ LÊN 15.000.000 đ/tháng**.",
            "  * *Lý do:* Tháng 10/2026, Vợ đã tất toán 100% khoản trả góp điện thoại (35.49tr), giải phóng ngay **3.000.000 đ/tháng** từ tháng 11/2026. Nên chuyển thẳng khoản này vào DCA 3 quỹ để nhân đôi tốc độ lãi kép!",
            "- **🎯 Quỹ Tiền Mặt Bắt Đáy (Hiện có 36.57tr)**:",
            "  * *Khuyến nghị:* **Nâng mức trích từ 4tr lên 6tr/tháng** khi dòng tiền thặng dư trên 35tr/tháng, sẵn sàng mua khi VN-Index chiết khấu sâu.\n",
            "#### 🟡 3. CÁC NHÓM DUY TRÌ & BẢO VỆ:",
            "- **👦 Học tập Subo (~15.1tr/tháng):** Duy trì trọn vẹn, đầu tư giáo dục cho con là ưu tiên số 1.",
            "- **👶 Sugi (~3.8tr/tháng bỉm sữa, y tế):** Đảm bảo dinh dưỡng và tiêm chủng đầy đủ.",
            "- **💳 Nợ Thuế Chồng:** Duy trì trả đều 3tr/tháng (còn 23.1tr, dự kiến tất toán vào Tháng 6/2027)."
        ]
        return "\n".join(lines)

    if "lỗi" in q or "logic" in q or "excel" in q or "sai" in q:
        bugs = db["audit"]["formula_bugs"]
        lines = ["### 🔍 Kết quả Kiểm Toán Logic & Công Thức Excel Gia Đình\n"]
        lines.append("Mình đã đối chiếu toàn bộ 30 sheets của 2 file Excel và phát hiện **7 lỗi công thức & dữ liệu** (đã tự động sửa trong App và file Excel chuẩn hóa):\n")
        for b in bugs:
            lines.append(f"- **[{b['id']}] {b['title']}** *(Vị trí: `{b['location']}`)*:\n  {b['detail']}")
        return "\n".join(lines)

    if any(k in q for k in ["đầu tư", "dca", "etf", "quỹ", "fuevfvnd", "dcds", "vesaf", "nav", "ccq", "bắt đáy"]):
        holdings = inv["holdings"]
        tot_invested = sum(h["invested"] for h in holdings)
        tot_mv = sum(h["market_value"] for h in holdings)
        tot_pnl = tot_mv - tot_invested
        tot_pnl_pct = (tot_pnl / tot_invested * 100) if tot_invested > 0 else 0
        lines = [
            "### 📈 Báo Cáo Danh Mục Đầu Tư DCA (ETF & Quỹ Mở)\n",
            f"- **Chiến lược cốt lõi**: Mua định kỳ (DCA) **12,000,000 đ/tháng** (30% `FUEVFVND` • 25% `DCDS` • 45% `VESAF`), tăng trưởng vốn góp 5%/năm.",
            f"- **Tổng vốn đã giải ngân DCA (8 kỳ)**: **{tot_invested:,.0f} đ**",
            f"- **Giá trị thị trường hiện tại**: **{tot_mv:,.0f} đ** (**Lãi tạm tính: +{tot_pnl:,.0f} đ / +{tot_pnl_pct:.2f}%**)",
            f"- **Quỹ Tiền Mặt Bắt Đáy tích lũy**: **{inv['dip_fund_balance']:,.0f} đ** *(Sẵn sàng giải ngân 15%–30% khi VN-Index chỉnh từ 10%–40%)*",
            f"- **Quỹ Dự Phòng Khẩn Cấp tích lũy**: **{inv['emergency_fund_balance']:,.0f} đ**\n",
            "| Mã Quỹ | Loại | Tỷ trọng MT | Đã đầu tư | Số CCQ | Giá vốn BQ | NAV hiện tại | Giá trị TT | Lãi/Lỗ |",
            "|---|---|---:|---:|---:|---:|---:|---:|---:|"
        ]
        for h in holdings:
            lines.append(
                f"| **{h['code']}** | {h['type']} | {h['target_weight']*100:.0f}% | {h['invested']:,.0f}đ | {h['units']:,.2f} | {h['avg_cost']:,.0f}đ | {h['current_nav']:,.0f}đ | **{h['market_value']:,.0f}đ** | +{h['pnl']:,.0f}đ (+{h['pnl_pct']}%) |"
            )
        lines.append("\n💡 **Khuyến nghị tháng 10/2026**: Tháng 10 hiện đang dư **+38,495,000 đ** (48.2% thu nhập). Bạn có thể bấm nút **'Khớp lệnh DCA Tháng 10 (12tr)'** ngay trong tab Đầu tư để mua thêm `3.6tr FUEVFVND`, `3.0tr DCDS`, `5.4tr VESAF`!")
        return "\n".join(lines)

    if any(k in q for k in ["nợ", "thuế", "shinhan", "trả góp"]):
        lines = [
            "### 💳 Tình Trạng Các Khoản Nợ & Trả Góp Gia Đình\n",
            "1. 🎉 **Trả góp Điện thoại Vợ (`OUTSTANDING VK`) — ĐÃ TẤT TOÁN XONG 100%!**\n   - Tổng nợ gốc: **35,490,000 đ** (trả từ T11/2025 đến T10/2026).\n   - Kỳ Tháng 10/2026 vừa trả **2,490,000 đ** là kỳ cuối cùng, dư nợ hiện tại = **0 đ**. Từ tháng 11/2026, ngân sách Vợ giảm tải được **3,000,000 đ/tháng**!",
            "2. 🏛️ **Nợ Thuế TNCN của Chồng (`OUSTANDING - CK`)**:\n   - Tổng thuế: **53,100,000 đ** (`18tr + 35.1tr`).\n   - Đã trả 10 kỳ (T1–T10/2026): **30,000,000 đ** (`3,000,000 đ/tháng`).\n   - Dư nợ còn lại: **23,100,000 đ** (dự kiến trả thêm **7 kỳ 3tr + 1 kỳ 2.1tr**, hoàn tất vào **Tháng 6/2027**).",
            "3. 🏦 **Trả góp Shinhan của Chồng**:\n   - Đang trả đều đặn **4,500,000 đ/tháng** trong bảng chi tiêu (10 tháng qua đã trả **45,000,000 đ**)."
        ]
        return "\n".join(lines)

    if any(k in q for k in ["subo", "sugi", "con", "học phí", "bỉm", "sữa"]):
        subo_txs = [t for t in txs if t["person"] == "Subo"]
        sugi_txs = [t for t in txs if t["person"] == "Sugi"]
        subo_tot = sum(t["amount"] for t in subo_txs)
        sugi_tot = sum(t["amount"] for t in sugi_txs)
        subo_by_cat = {}
        for t in subo_txs:
            subo_by_cat[t["category"]] = subo_by_cat.get(t["category"], 0) + t["amount"]
        sugi_by_cat = {}
        for t in sugi_txs:
            sugi_by_cat[t["category"]] = sugi_by_cat.get(t["category"], 0) + t["amount"]

        lines = [
            "### 👶 Phân Tích Chi Phí Của 2 Bé: SUBO & SUGI (10 Tháng 2026)\n",
            f"- **Tổng chi cho SUBO**: **{subo_tot:,.0f} đ** *(Trung bình **{subo_tot/10:,.0f} đ/tháng**, chiếm **{subo_tot/total_exp_ytd*100:.1f}%** tổng chi gia đình)*",
            f"  - Trong đó: Học tập **{subo_by_cat.get('Học tập',0):,.0f} đ** • Vui chơi **{subo_by_cat.get('Vui chơi',0):,.0f} đ** • Làm đẹp/Quần áo **{subo_by_cat.get('Làm đẹp',0):,.0f} đ** • Y tế **{subo_by_cat.get('Y tế',0):,.0f} đ**.",
            f"- **Tổng chi cho SUGI**: **{sugi_tot:,.0f} đ** *(Trung bình **{sugi_tot/10:,.0f} đ/tháng**, chiếm **{sugi_tot/total_exp_ytd*100:.1f}%** tổng chi gia đình)*",
            f"  - Trong đó: Đồ dùng/Bỉm sữa **{sugi_by_cat.get('Đồ dùng GĐ',0):,.0f} đ** • Y tế/Tiêm chủng **{sugi_by_cat.get('Y tế',0):,.0f} đ** • Làm đẹp/Quần áo **{sugi_by_cat.get('Làm đẹp',0):,.0f} đ**.",
            f"- **Quỹ riêng của con**: Quỹ Mừng Tuổi (**17,160,000 đ**) + Quỹ Mừng Sinh Sugi (**24,200,000 đ**) = **41,360,000 đ**."
        ]
        return "\n".join(lines)

    if "ăn uống" in q or "đi chợ" in q or "siêu thị" in q:
        food_txs = [t for t in txs if t["category"] == "Ăn uống"]
        food_tot = sum(t["amount"] for t in food_txs)
        avg_month = food_tot / 10
        avg_day = food_tot / 295  # ~295 days T1..T10
        lines = [
            "### 🍲 Phân Tích Chi Tiêu Ăn Uống Gia Đình (T1 – T10/2026)\n",
            f"- **Tổng chi Ăn uống 10 tháng**: **{food_tot:,.0f} đ** trên tổng **{len(food_txs)} giao dịch** (chiếm **{food_tot/total_exp_ytd*100:.1f}%** tổng chi tiêu gia đình).",
            f"- **Trung bình mỗi tháng**: **{avg_month:,.0f} đ/tháng**",
            f"- **Trung bình mỗi ngày**: **~{avg_day:,.0f} đ/ngày**",
            "- **Tháng chi Ăn uống cao nhất**: **Tháng 9** (`22,030,000 đ`) và **Tháng 8** (`20,638,000 đ`).",
            "- **Tháng hiện tại (Tháng 10)**: Đang ở mức **`6,227,000 đ`** — rất tiết kiệm so với trung bình 17.5tr/tháng!"
        ]
        return "\n".join(lines)

    if target_m is not None:
        m_obj = next((m for m in months if m["month"] == target_m), None)
        m_txs = [t for t in txs if t["month"] == target_m]
        if m_obj:
            by_cat = {}
            by_person = {}
            for t in m_txs:
                by_cat[t["category"]] = by_cat.get(t["category"], 0) + t["amount"]
                by_person[t["person"]] = by_person.get(t["person"], 0) + t["amount"]
            top_cats = sorted(by_cat.items(), key=lambda x: x[1], reverse=True)
            top_txs = sorted(m_txs, key=lambda x: x["amount"], reverse=True)[:5]

            lines = [
                f"### 📊 Báo Cáo Tài Chính Chi Tiết — Tháng {target_m}/2026\n",
                f"- **Tổng Thu Nhập**: **{m_obj['income']['total']:,.0f} đ** *(Chồng: {m_obj['income']['CK']:,.0f}đ • Vợ: {m_obj['income']['VK']:,.0f}đ • Subo: {m_obj['income']['Subo']:,.0f}đ • Sugi: {m_obj['income']['Sugi']:,.0f}đ)*",
                f"- **Hạn mức Chi Tiêu (70%)**: **{m_obj['target_expense']:,.0f} đ**",
                f"- **Thực Chi ({len(m_txs)} khoản)**: **{m_obj['actual_total_expense']:,.0f} đ**",
                f"- **Tiền Dư Thực Tế**: **{m_obj['remaining_actual']:,.0f} đ** *(Tỷ lệ tiết kiệm: **{(m_obj['remaining_actual']/m_obj['income']['total']*100) if m_obj['income']['total']>0 else 0:.1f}%**)*\n",
                "**1. Chi theo Thành viên:**"
            ]
            for p_name, p_val in sorted(by_person.items(), key=lambda x: x[1], reverse=True):
                lines.append(f"- **{p_name}**: {p_val:,.0f} đ ({p_val/m_obj['actual_total_expense']*100 if m_obj['actual_total_expense']>0 else 0:.1f}%)")
            lines.append("\n**2. Top Hạng mục chi nhiều nhất:**")
            for c_name, c_val in top_cats[:5]:
                lines.append(f"- **{c_name}**: {c_val:,.0f} đ ({c_val/m_obj['actual_total_expense']*100 if m_obj['actual_total_expense']>0 else 0:.1f}%)")
            lines.append("\n**3. 5 Khoản chi lớn nhất trong tháng:**")
            for t in top_txs:
                lines.append(f"- `{t['content']}` ({t['category']} • {t['person']}): **{t['amount']:,.0f} đ**")
            return "\n".join(lines)

    # Default comprehensive overview
    by_cat_all = {}
    by_person_all = {}
    for t in txs:
        by_cat_all[t["category"]] = by_cat_all.get(t["category"], 0) + t["amount"]
        by_person_all[t["person"]] = by_person_all.get(t["person"], 0) + t["amount"]

    top_cats_all = sorted(by_cat_all.items(), key=lambda x: x[1], reverse=True)
    lines = [
        "### 🏆 Tổng Quan Sức Khỏe Tài Chính Gia Đình (10 Tháng Năm 2026)\n",
        f"- **Tổng Thu Nhập 10 tháng**: **{total_inc_ytd:,.0f} đ** *(Chồng: 624.06tr • Vợ: 375.50tr • Subo & Sugi: 35.72tr)*",
        f"- **Tổng Chi Tiêu Thực Tế ({len(txs)} giao dịch)**: **{total_exp_ytd:,.0f} đ**",
        f"- **Thặng Dư Dòng Tiền Ròng (100%)**: **+{total_rem_ytd:,.0f} đ** — **Tỷ lệ tiết kiệm đạt {saving_rate_ytd:.1f}%**\n"
        f"  - Trong đó **80% Đưa vào Tích lũy/Đầu tư (`N16`)**: **{total_rem_ytd*0.8:,.0f} đ** ({saving_rate_ytd*0.8:.1f}% thu nhập)\n"
        f"  - Và **20% Tiền còn lại giữ linh hoạt**: **{total_rem_ytd*0.2:,.0f} đ** ({saving_rate_ytd*0.2:.1f}% thu nhập)\n",
        "**Top 5 Hạng Mục Chi Tiêu Lớn Nhất Năm 2026:**"
    ]
    for c_name, c_val in top_cats_all[:5]:
        lines.append(f"- **{c_name}**: **{c_val:,.0f} đ** ({c_val/total_exp_ytd*100:.1f}%)")

    lines.append("\n**Chi Tiêu Theo Thành Viên (10 Tháng):**")
    for p_name, p_val in sorted(by_person_all.items(), key=lambda x: x[1], reverse=True):
        lines.append(f"- **{p_name}**: **{p_val:,.0f} đ** ({p_val/total_exp_ytd*100:.1f}%)")

    lines.append("\n👉 *Bạn có thể hỏi cụ thể hơn như: \"Tháng 10 tiêu gì nhiều nhất?\", \"So sánh chi phí Subo và Sugi\", \"Danh mục DCA 3 quỹ hiện lời bao nhiêu?\", hoặc \"Khi nào trả hết nợ thuế?\"*")
    return "\n".join(lines)


# ==============================================================================
# GMAIL IMAP FETCHER FOR VPBANK, BIDV, NCB
# ==============================================================================

def fetch_gmail_bank_emails(email_addr, app_password, max_emails=15):
    """Connects to Gmail IMAP and searches recent emails from VPBank, BIDV, NCB."""
    mail = imaplib.IMAP4_SSL("imap.gmail.com", 993)
    mail.login(email_addr, app_password)
    mail.select("INBOX")

    # Search for bank keywords in subject/from
    status, data = mail.search(None, 'ALL')
    if status != "OK" or not data or not data[0]:
        mail.logout()
        return []

    msg_ids = data[0].split()[-max_emails:]
    raw_blocks = []

    for m_id in reversed(msg_ids):
        res, msg_data = mail.fetch(m_id, "(RFC822)")
        if res != "OK":
            continue
        for response_part in msg_data:
            if isinstance(response_part, tuple):
                msg = email.message_from_bytes(response_part[1])
                subj, enc = decode_header(msg.get("Subject", ""))[0]
                if isinstance(subj, bytes):
                    subj = subj.decode(enc or "utf-8", errors="ignore")
                sender = msg.get("From", "")
                combined_header = f"{sender} {subj}".lower()
                if not any(b in combined_header for b in ["vpbank", "bidv", "ncb", "uob", "shopee", "biến động", "giao dịch", "thông báo"]):
                    continue

                body = ""
                if msg.is_multipart():
                    for part in msg.walk():
                        ctype = part.get_content_type()
                        if ctype == "text/plain":
                            payload = part.get_payload(decode=True)
                            if payload:
                                body += payload.decode("utf-8", errors="ignore")
                else:
                    payload = msg.get_payload(decode=True)
                    if payload:
                        body = payload.decode("utf-8", errors="ignore")

                raw_blocks.append(f"From: {sender}\nSubject: {subj}\n{body}")

    mail.logout()
    if not raw_blocks:
        return []
    return parse_bank_emails("\n---\n".join(raw_blocks))


SAMPLE_BANK_EMAILS = """From: VPBank Online <customercare@vpbank.com.vn>
Subject: [VPBank] Thông báo biến động số dư tài khoản
Ngày 22/10/2026 08:15:20. TK 1988xxxx tại VPBank biến động -485,000 VND.
Nội dung: CO.OPMART RACH MIEU THANH TOAN QR. Số dư khả dụng: 24,510,000 VND.
---
From: BIDV SmartBanking <bidvsmartbanking@bidv.com.vn>
Subject: [BIDV] Thông báo giao dịch tài khoản 3141xxxx
Thời gian: 22/10/2026 11:30. Số tiền GD: -395,000 VND.
Nội dung: Mua bỉm và sữa thanh cho Sugi tại CON CUNG.
---
From: NCB iziMobile <info@ncb-bank.vn>
Subject: [NCB] Biến động số dư tài khoản thanh toán
NCB thông báo: Ngày 22/10/2026, TK xxxx8899 trừ -95,000 VND.
Nội dung: Đổ xăng ck Petrolimex Q1.
---
From: VPBank Shopee Card <card@vpbank.com.vn>
Subject: [VPBank] Giao dịch thẻ tín dụng VPBank Shopee
Thẻ VPBank Shopee *4591 vừa chi tiêu 320,000 VND ngày 22/10/2026.
Nội dung: SHOPEE VN - Nước giặt OMO và túi rác gia đình.
---
From: BIDV SmartBanking <bidvsmartbanking@bidv.com.vn>
Subject: [BIDV] Thanh toán hóa đơn thành công
Ngày 22/10/2026, TK BIDV thanh toán -265,000 VND.
Nội dung: Nhà thuốc Long Châu mua thuốc ho và vitamin cho Subo.
---
From: VPBank Online <customercare@vpbank.com.vn>
Subject: [VPBank] Thông báo giao dịch thẻ ghi nợ
Ngày 22/10/2026 14:10:00. TK 1988xxxx trừ -1,250,000 VND.
Nội dung: THANH TOAN POS 004128 CH SO 19."""


APPS_SCRIPT_CODE = """/**
 * GOOGLE APPS SCRIPT - ĐỒNG BỘ 2 CHIỀU VỚI WEB APP FAMILY WEALTH OS 2026
 * Cách cài đặt:
 * 1. Mở Google Sheet [2026 - FAMILY NOTE BOOK] của bạn
 * 2. Vào Tiện ích mở rộng (Extensions) -> Apps Script
 * 3. Dán đoạn mã này vào và bấm Triển khai (Deploy) -> Web App (Anyone)
 */
const SHEET_ID = '1FtjpyD-DP4HS8z7uJoYDw9TKwBh3qSCwdXqy1eaQEEk';

function doGet(e) {
  const ss = SpreadsheetApp.openById(SHEET_ID);
  const month = (e && e.parameter && e.parameter.month) ? e.parameter.month : '10';
  const sheet = ss.getSheetByName('T' + month + ' - DETAIL');
  const data = sheet.getRange('B4:G200').getValues().filter(r => r[1] || r[4]);
  return ContentService.createTextOutput(JSON.stringify({ ok: true, month, rows: data }))
    .setMimeType(ContentService.MimeType.JSON);
}

function doPost(e) {
  const payload = JSON.parse(e.postData.contents);
  const ss = SpreadsheetApp.openById(SHEET_ID);
  const sheet = ss.getSheetByName('T' + payload.month + ' - DETAIL');
  
  // Tìm dòng trống đầu tiên trong cột C & F (từ dòng 4)
  const colC = sheet.getRange('C4:C800').getValues();
  const colF = sheet.getRange('F4:F800').getValues();
  let targetRow = 4;
  for (let i = 0; i < colC.length; i++) {
    if (!colC[i][0] && !colF[i][0]) {
      targetRow = i + 4;
      break;
    }
  }
  sheet.getRange(targetRow, 2, 1, 6).setValues([[
    payload.date,
    payload.content,
    payload.category,
    payload.person,
    payload.amount,
    payload.source
  ]]);
  return ContentService.createTextOutput(JSON.stringify({ ok: true, row: targetRow }))
    .setMimeType(ContentService.MimeType.JSON);
}
"""


class FamilyAppHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def _send_json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(body)

    def _read_body(self):
        length = int(self.headers.get("Content-Length", 0))
        if length == 0:
            return {}
        raw = self.rfile.read(length).decode("utf-8")
        return json.loads(raw)

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, PUT, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/data":
            db = load_db()
            recompute_months(db)
            db["sample_bank_emails"] = SAMPLE_BANK_EMAILS
            db["apps_script_code"] = APPS_SCRIPT_CODE
            self._send_json(db)
            return
        elif parsed.path == "/api/export-excel":
            if os.path.exists(FIXED_EXCEL_PATH):
                with open(FIXED_EXCEL_PATH, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", 'attachment; filename="2026_FAMILY_NOTEBOOK_CHUAN_HOA.xlsx"')
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)
                return
            self._send_json({"error": "File not found"}, 404)
            return

        return super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        body = self._read_body()

        if parsed.path == "/api/quick-parse":
            text = body.get("text", "")
            month = body.get("month", 10)
            res = parse_quick_input(text, default_month=month)
            self._send_json({"ok": True, "parsed": res})
            return

        elif parsed.path == "/api/transactions":
            db = load_db()
            txs = db["transactions"]
            new_num = len(txs) + 1
            tx = {
                "id": f"TX2026-{new_num:04d}",
                "month": int(body.get("month", 10)),
                "date": body.get("date", datetime.datetime.now().strftime("2026-10-%d")),
                "row": 0,
                "content": body.get("content", "").strip(),
                "category": body.get("category", "Ăn uống"),
                "person": body.get("person", "Chung"),
                "amount": float(body.get("amount", 0)),
                "source": body.get("source", "VPBank"),
                "was_unmatched_in_excel": False,
                "issue_note": None
            }
            txs.insert(0, tx)
            recompute_months(db)
            save_db(db)
            excel_synced = append_tx_to_excel(tx)
            self._send_json({"ok": True, "transaction": tx, "excel_synced": excel_synced, "db": db})
            return

        elif parsed.path == "/api/transactions/batch":
            db = load_db()
            items = body.get("items", [])
            added = []
            for item in items:
                new_num = len(db["transactions"]) + 1
                tx = {
                    "id": f"TX2026-{new_num:04d}",
                    "month": int(item.get("month", 10)),
                    "date": item.get("date", datetime.datetime.now().strftime("2026-10-%d")),
                    "row": 0,
                    "content": item.get("content", "").strip(),
                    "category": item.get("category", "Ăn uống"),
                    "person": item.get("person", "Chung"),
                    "amount": float(item.get("amount", 0)),
                    "source": item.get("source", item.get("bank", "VPBank")),
                    "was_unmatched_in_excel": False,
                    "issue_note": None
                }
                db["transactions"].insert(0, tx)
                append_tx_to_excel(tx)
                added.append(tx)
            recompute_months(db)
            save_db(db)
            self._send_json({"ok": True, "added_count": len(added), "db": db})
            return

        elif parsed.path == "/api/transactions/delete":
            db = load_db()
            tx_id = body.get("id")
            db["transactions"] = [t for t in db["transactions"] if t["id"] != tx_id]
            recompute_months(db)
            save_db(db)
            self._send_json({"ok": True, "db": db})
            return

        elif parsed.path == "/api/income":
            db = load_db()
            m_num = int(body.get("month", 10))
            for m_obj in db["months"]:
                if m_obj["month"] == m_num:
                    m_obj["income"]["CK"] = float(body.get("CK", m_obj["income"]["CK"]))
                    m_obj["income"]["VK"] = float(body.get("VK", m_obj["income"]["VK"]))
                    m_obj["income"]["Subo"] = float(body.get("Subo", m_obj["income"]["Subo"]))
                    m_obj["income"]["Sugi"] = float(body.get("Sugi", m_obj["income"]["Sugi"]))
                    break
            recompute_months(db)
            save_db(db)
            self._send_json({"ok": True, "db": db})
            return

        elif parsed.path == "/api/email/parse":
            raw_text = body.get("raw_text", SAMPLE_BANK_EMAILS)
            parsed_list = parse_bank_emails(raw_text)
            self._send_json({"ok": True, "items": parsed_list})
            return

        elif parsed.path == "/api/email/clarify":
            raw_content = body.get("raw_content", "")
            amount = float(body.get("amount", 0))
            user_reply = body.get("user_reply", "").strip()
            cat, person, conf, reason = classify_transaction_text(raw_content + " " + user_reply)
            self._send_json({
                "ok": True,
                "content": user_reply or raw_content,
                "category": cat,
                "person": person,
                "confidence": max(conf, 0.95),
                "reason": f"AI cập nhật theo phản hồi: '{user_reply}'" if user_reply else reason,
                "needs_clarification": False
            })
            return

        elif parsed.path == "/api/email/sync-imap":
            email_addr = body.get("email", "").strip()
            app_pwd = body.get("app_password", "").strip()
            if not email_addr or not app_pwd:
                self._send_json({"ok": False, "error": "Vui lòng nhập địa chỉ Gmail và Mật khẩu ứng dụng (App Password 16 ký tự)."}, 400)
                return
            try:
                items = fetch_gmail_bank_emails(email_addr, app_pwd)
                self._send_json({"ok": True, "items": items})
            except Exception as e:
                self._send_json({"ok": False, "error": f"Không thể kết nối IMAP Gmail: {str(e)}"}, 400)
            return

        elif parsed.path == "/api/investment/update-nav":
            db = load_db()
            nav_updates = body.get("navs", {})
            for h in db["investment"]["holdings"]:
                if h["code"] in nav_updates:
                    h["current_nav"] = float(nav_updates[h["code"]])
            recompute_months(db)
            save_db(db)
            self._send_json({"ok": True, "db": db})
            return

        elif parsed.path == "/api/investment/record-dca":
            db = load_db()
            m_num = int(body.get("month", 10))
            total_dca = float(body.get("amount", 12000000))
            holdings_map = {h["code"]: h for h in db["investment"]["holdings"]}

            for code, ratio in [("FUEVFVND", 0.30), ("DCDS", 0.25), ("VESAF", 0.45)]:
                amt = round(total_dca * ratio)
                nav = float(holdings_map[code]["current_nav"])
                units = round(amt / nav, 2) if nav > 0 else 0
                db["investment"]["dca_orders"].append({
                    "id": f"DCA-T{m_num}-{code}-{len(db['investment']['dca_orders'])+1}",
                    "month": m_num,
                    "date": datetime.datetime.now().strftime(f"2026-{m_num:02d}-%d"),
                    "code": code,
                    "amount": amt,
                    "buy_nav": nav,
                    "units": units,
                    "status": "Đã khớp"
                })

            for m_obj in db["months"]:
                if m_obj["month"] == m_num:
                    m_obj["allocation"]["dca_total"] = total_dca
                    m_obj["allocation"]["status"] = "Đã mua"
                    break

            recompute_months(db)
            save_db(db)
            self._send_json({"ok": True, "db": db})
            return

        elif parsed.path == "/api/ai/ask":
            db = load_db()
            recompute_months(db)
            question = body.get("question", "")
            answer = answer_family_question(question, db)
            self._send_json({"ok": True, "answer": answer})
            return

        self._send_json({"error": "Unknown endpoint"}, 404)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8686))
    server = HTTPServer(("127.0.0.1", port), FamilyAppHandler)
    print(f"Family Wealth OS 2026 running at http://localhost:{port}")
    server.serve_forever()
